"""Imports a reviewed gvpn-translation-review.xlsx back into lib/i18n.

Usage (from apps/g_vpn):
    python3 tool/import_review.py ~/Downloads/gvpn-translation-review.xlsx

Only rows with something in the Corrected column change. A correction is
refused, and reported, when its {placeholders} differ from the English
sentence's, because the app would then show a raw {n} or drop a number.
Needs openpyxl (pip install openpyxl). Run `dart run slang` afterwards.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

I18N = Path(__file__).resolve().parent.parent / 'lib' / 'i18n'
TABS = {'Persian': 'fa', 'Russian': 'ru'}
HOLES = re.compile(r'\{[a-zA-Z]+\}')


def put(tree, key, value):
    parts = key.split('.')
    for p in parts[:-1]:
        tree = tree.setdefault(p, {})
    tree[parts[-1]] = value


def main(path):
    wb = load_workbook(path, data_only=True)
    for tab, code in TABS.items():
        if tab not in wb.sheetnames:
            continue
        file = I18N / f'{code}.json'
        raw = file.read_text(encoding='utf-8')
        data = json.loads(raw)
        changed, refused = 0, []
        for row in wb[tab].iter_rows(min_row=2, values_only=True):
            key, _, english, _, fixed = row[:5]
            if not key or not fixed or not str(fixed).strip():
                continue
            fixed = str(fixed).strip()
            if sorted(HOLES.findall(str(english))) != sorted(HOLES.findall(fixed)):
                refused.append(key)
                continue
            put(data, key, fixed)
            changed += 1
        file.write_text(json.dumps(data, ensure_ascii=False, indent=2) + ('\n' if raw.endswith('\n') else ''), encoding='utf-8')
        print(f'{tab}: {changed} corrected' + (f', {len(refused)} refused (placeholders differ): {", ".join(refused)}' if refused else ''))


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
